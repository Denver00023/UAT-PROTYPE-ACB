import streamlit as st
import pandas as pd

from io import BytesIO

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


# Clean values for consistent matching.
def clean_value(value):

    if pd.isna(value):
        return ""

    value = str(value).strip()

    if value.endswith(".0"):
        value = value[:-2]

    return value


# Normalize descriptions for client-file matching.
def normalize_description(value):

    value = clean_value(value)

    value = " ".join(
        value.upper().split()
    )

    return value


# Collect Client_HS_code from the third file.
def add_client_hs_code(
    shipment_df,
    client_df
):

    required_shipment_columns = [
        "Client_Internal_tracking",
        "Goods_Description"
    ]

    required_client_columns = [
        "Reliable_tracking",
        "Goods_Description",
        "HS_code"
    ]

    for column in required_shipment_columns:

        if column not in shipment_df.columns:

            raise ValueError(
                f"Second file must contain: {column}"
            )

    for column in required_client_columns:

        if column not in client_df.columns:

            raise ValueError(
                f"Third file must contain: {column}"
            )

    result_df = shipment_df.copy()

    client_hs_lookup = {}

    for _, row in client_df.iterrows():

        reliable_tracking = clean_value(
            row["Reliable_tracking"]
        )

        goods_description = normalize_description(
            row["Goods_Description"]
        )

        client_hs = clean_value(
            row["HS_code"]
        )

        if not reliable_tracking:
            continue

        if not goods_description:
            continue

        if not client_hs:
            continue

        lookup_key = (
            reliable_tracking,
            goods_description
        )

        client_hs_lookup[lookup_key] = client_hs

    result_df["Client_HS_code"] = ""

    result_df["_Client_HS_Status"] = (
        "No Client Match"
    )

    for index, row in result_df.iterrows():

        internal_tracking = clean_value(
            row["Client_Internal_tracking"]
        )

        goods_description = normalize_description(
            row["Goods_Description"]
        )

        lookup_key = (
            internal_tracking,
            goods_description
        )

        if lookup_key in client_hs_lookup:

            client_hs = client_hs_lookup[
                lookup_key
            ]

            result_df.at[
                index,
                "Client_HS_code"
            ] = client_hs

            result_df.at[
                index,
                "_Client_HS_Status"
            ] = "Matched"

    return result_df


# Use Client_HS_code to look up the first file master.
# The master file determines the final HS_code.
def apply_master_hs_adjustment(
    result_df,
    reference_df
):

    required_reference_columns = [
        "Client_HS_code",
        "Adjusted_HS_code"
    ]

    for column in required_reference_columns:

        if column not in reference_df.columns:

            raise ValueError(
                f"First file must contain: {column}"
            )

    if "HS_code" not in result_df.columns:

        raise ValueError(
            "Second file must contain: HS_code"
        )

    result_df = result_df.copy()

    result_df["_Original_HS_code"] = (
        result_df["HS_code"]
        .apply(clean_value)
    )

    result_df["_HS_Code_Status"] = (
        "No Client HS Match"
    )

    result_df["_HS_Code_Changed"] = False

    master_lookup = {}

    for _, row in reference_df.iterrows():

        client_hs = clean_value(
            row["Client_HS_code"]
        )

        adjusted_hs = clean_value(
            row["Adjusted_HS_code"]
        )

        if not client_hs:
            continue

        master_lookup[client_hs] = adjusted_hs

    for index, row in result_df.iterrows():

        client_hs = clean_value(
            row["Client_HS_code"]
        )

        original_hs = clean_value(
            row["HS_code"]
        )

        if not client_hs:

            result_df.at[
                index,
                "_HS_Code_Status"
            ] = "No Client HS Match"

            continue

        if client_hs not in master_lookup:

            result_df.at[
                index,
                "_HS_Code_Status"
            ] = "Client HS Not in Master"

            continue

        adjusted_hs = master_lookup[
            client_hs
        ]

        if not adjusted_hs:

            result_df.at[
                index,
                "_HS_Code_Status"
            ] = "Master Match / No Adjustment"

            continue

        if original_hs != adjusted_hs:

            result_df.at[
                index,
                "HS_code"
            ] = adjusted_hs

            result_df.at[
                index,
                "_HS_Code_Status"
            ] = "Client HS → Master Adjustment"

            result_df.at[
                index,
                "_HS_Code_Changed"
            ] = True

        else:

            result_df.at[
                index,
                "_HS_Code_Status"
            ] = "Master Adjustment / Same Code"

    return result_df


# Prepare the final export while preserving shipment columns.
def prepare_final_dataframe(
    result_df,
    original_columns
):

    export_columns = original_columns.copy()

    if "Client_HS_code" in export_columns:

        export_columns.remove(
            "Client_HS_code"
        )

    if "HS_code" in export_columns:

        hs_index = export_columns.index(
            "HS_code"
        )

        export_columns.insert(
            hs_index,
            "Client_HS_code"
        )

    else:

        export_columns.append(
            "Client_HS_code"
        )

    export_df = result_df[
        export_columns
    ].copy()

    if "Client_HS_code" in export_df.columns:

        export_df["Client_HS_code"] = (
            export_df["Client_HS_code"]
            .fillna("")
            .astype(str)
        )

    if "HS_code" in export_df.columns:

        export_df["HS_code"] = (
            export_df["HS_code"]
            .fillna("")
            .astype(str)
        )

    return export_df


# Create and format the final Excel workbook.
def create_excel(
    result_df,
    original_columns
):

    export_df = prepare_final_dataframe(
        result_df,
        original_columns
    )

    output = BytesIO()

    with pd.ExcelWriter(
        output,
        engine="openpyxl"
    ) as writer:

        export_df.to_excel(
            writer,
            index=False,
            sheet_name="HS Validation"
        )

    output.seek(0)

    workbook = load_workbook(
        output
    )

    worksheet = workbook[
        "HS Validation"
    ]

    headers = {
        cell.value: cell.column
        for cell in worksheet[1]
    }

    hs_column = headers.get(
        "HS_code"
    )

    client_hs_column = headers.get(
        "Client_HS_code"
    )

    changed_fill = PatternFill(
        fill_type="solid",
        fgColor="C6EFCE"
    )

    changed_font = Font(
        bold=True,
        color="006100"
    )

    client_fill = PatternFill(
        fill_type="solid",
        fgColor="FFF2CC"
    )

    client_font = Font(
        bold=True,
        color="9C6500"
    )

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78"
    )

    header_font = Font(
        bold=True,
        color="FFFFFF"
    )

    # Highlight HS codes adjusted through the master file.
    if hs_column:

        for excel_row, changed in enumerate(
            result_df["_HS_Code_Changed"],
            start=2
        ):

            if changed:

                cell = worksheet.cell(
                    row=excel_row,
                    column=hs_column
                )

                cell.fill = changed_fill
                cell.font = changed_font

    # Highlight Client_HS_code values collected from the third file.
    if client_hs_column:

        for excel_row, value in enumerate(
            result_df["Client_HS_code"],
            start=2
        ):

            if clean_value(value):

                cell = worksheet.cell(
                    row=excel_row,
                    column=client_hs_column
                )

                cell.fill = client_fill
                cell.font = client_font

    # Format the Excel header.
    for cell in worksheet[1]:

        cell.fill = header_fill

        cell.font = header_font

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

    worksheet.auto_filter.ref = (
        f"A2:{get_column_letter(worksheet.max_column)}{worksheet.max_row}"
    )

    # Automatically size Excel columns.
    for column_cells in worksheet.columns:

        max_length = 0

        column_letter = get_column_letter(
            column_cells[0].column
        )

        for cell in column_cells:

            try:

                max_length = max(
                    max_length,
                    len(str(cell.value))
                )

            except Exception:
                pass

        worksheet.column_dimensions[
            column_letter
        ].width = min(
            max_length + 2,
            50
        )

    final_output = BytesIO()

    workbook.save(
        final_output
    )

    final_output.seek(0)

    return final_output


# Run the Streamlit application.
def run():

    st.title(
        "🔎 APC HS Code Validation"
    )

    st.markdown(
        """
        **First File — HS Master File**

        `Client_HS_code` = Client HS code used for master lookup  
        `Adjusted_HS_code` = Correct replacement HS code

        **Second File — Shipment File**

        `HS_code` = Original shipment HS code  
        `Client_Internal_tracking` = Tracking used for client matching

        **Third File — Client File**

        `Reliable_tracking` = Client tracking number  
        `Goods_Description` = Product description  
        `HS_code` = Client HS code

        The Client HS code is collected from the Third File
        and then looked up in the First File Master.

        If the Client HS code is found in the Master,
        its `Adjusted_HS_code` becomes the final `HS_code`.

        If the Client HS code is not found in the Master,
        the original shipment `HS_code` remains unchanged.
        """
    )

    st.divider()

    col1, col2, col3 = st.columns(3)

    with col1:

        st.subheader(
            "1️⃣ HS Master File"
        )

        first_file = st.file_uploader(
            "Upload First File",
            type=[
                "xlsx",
                "xls"
            ],
            key="apc_hs_first_file"
        )

        st.caption(
            "Client_HS_code + Adjusted_HS_code"
        )

    with col2:

        st.subheader(
            "2️⃣ Shipment File"
        )

        second_file = st.file_uploader(
            "Upload Second File",
            type=[
                "xlsx",
                "xls"
            ],
            key="apc_hs_second_file"
        )

        st.caption(
            "Main shipment file"
        )

    with col3:

        st.subheader(
            "3️⃣ Client File"
        )

        third_file = st.file_uploader(
            "Upload Client File",
            type=[
                "xlsx",
                "xls"
            ],
            key="apc_hs_third_file"
        )

        st.caption(
            "Client HS information"
        )

    st.divider()

    if (
        first_file
        and second_file
        and third_file
    ):

        if st.button(
            "🚀 Validate HS Codes",
            type="primary",
            use_container_width=True
        ):

            try:

                # Read the master file.
                with st.spinner(
                    "Reading HS master file..."
                ):

                    reference_df = pd.read_excel(
                        first_file,
                        dtype=str
                    )

                    reference_df.columns = (
                        reference_df.columns
                        .str.strip()
                    )

                # Read the shipment file.
                with st.spinner(
                    "Reading shipment file..."
                ):

                    shipment_df = pd.read_excel(
                        second_file,
                        dtype=str,
                        
                    )

                    shipment_df.columns = (
                        shipment_df.columns
                        .str.strip()
                    )

                # Read the client file.
                with st.spinner(
                    "Reading client file..."
                ):

                    client_df = pd.read_excel(
                        third_file,
                        dtype=str
                    )

                    client_df.columns = (
                        client_df.columns
                        .str.strip()
                    )

                original_columns = (
                    shipment_df.columns.tolist()
                )

                # Collect Client_HS_code from the third file.
                with st.spinner(
                    "Matching client information..."
                ):

                    result_df = add_client_hs_code(
                        shipment_df,
                        client_df
                    )

                # Use Client_HS_code to check the master file.
                with st.spinner(
                    "Checking Client HS codes against master..."
                ):

                    result_df = apply_master_hs_adjustment(
                        result_df,
                        reference_df
                    )

                # Create the final Excel file.
                with st.spinner(
                    "Creating final Excel file..."
                ):

                    excel_file = create_excel(
                        result_df,
                        original_columns
                    )

                adjusted_count = int(
                    result_df[
                        "_HS_Code_Changed"
                    ].sum()
                )

                client_match_count = int(
                    (
                        result_df[
                            "_Client_HS_Status"
                        ]
                        == "Matched"
                    ).sum()
                )

                client_no_match_count = int(
                    (
                        result_df[
                            "_Client_HS_Status"
                        ]
                        == "No Client Match"
                    ).sum()
                )

                master_match_count = int(
                    (
                        ~result_df[
                            "_HS_Code_Status"
                        ].isin([
                            "No Client HS Match",
                            "Client HS Not in Master"
                        ])
                    ).sum()
                )

                client_hs_not_in_master_count = int(
                    (
                        result_df[
                            "_HS_Code_Status"
                        ]
                        == "Client HS Not in Master"
                    ).sum()
                )

                total_rows = len(
                    result_df
                )

                st.success(
                    "HS code validation completed successfully."
                )

                st.subheader(
                    "📊 Validation Summary"
                )

                c1, c2, c3, c4 = st.columns(4)

                c1.metric(
                    "Total Rows",
                    total_rows
                )

                c2.metric(
                    "HS Codes Adjusted",
                    adjusted_count
                )

                c3.metric(
                    "Client HS Matched",
                    client_match_count
                )

                c4.metric(
                    "Client HS Not Matched",
                    client_no_match_count
                )

                st.caption(
                    f"Master HS matches: {master_match_count}  |  "
                    f"Client HS codes not in master: "
                    f"{client_hs_not_in_master_count}"
                )

                # Show the main validation result.
                st.divider()

                st.subheader(
                    "🔎 DataFrame Preview"
                )

                preview_columns = [
                    "Client_Internal_tracking",
                    "Goods_Description",
                    "_Original_HS_code",
                    "Client_HS_code",
                    "HS_code",
                    "_HS_Code_Status",
                    "_Client_HS_Status"
                ]

                available_preview_columns = [
                    column
                    for column in preview_columns
                    if column in result_df.columns
                ]

                preview_df = result_df[
                    available_preview_columns
                ].copy()

                preview_column_names = {

                    "Client_Internal_tracking":
                        "Client Internal Tracking",

                    "Goods_Description":
                        "Goods Description",

                    "_Original_HS_code":
                        "Original HS_code",

                    "Client_HS_code":
                        "Client_HS_code",

                    "HS_code":
                        "Final HS_code",

                    "_HS_Code_Status":
                        "HS Validation",

                    "_Client_HS_Status":
                        "Client HS Match"
                }

                preview_df = preview_df.rename(
                    columns=preview_column_names
                )

                st.dataframe(
                    preview_df,
                    use_container_width=True,
                    hide_index=True
                )

                # Show HS codes changed by the master.
                if adjusted_count > 0:

                    st.divider()

                    st.subheader(
                        "🔄 Adjusted HS Codes"
                    )

                    changed_df = result_df[
                        result_df[
                            "_HS_Code_Changed"
                        ]
                    ][
                        [
                            "_Original_HS_code",
                            "Client_HS_code",
                            "HS_code"
                        ]
                    ].copy()

                    changed_df.columns = [
                        "Original HS_code",
                        "Client_HS_code",
                        "Adjusted HS_code"
                    ]

                    st.dataframe(
                        changed_df,
                        use_container_width=True,
                        hide_index=True
                    )

                # Show Client HS codes collected from the third file.
                if client_match_count > 0:

                    st.divider()

                    st.subheader(
                        "👤 Client HS Code Matches"
                    )

                    client_preview = result_df[
                        result_df[
                            "_Client_HS_Status"
                        ]
                        == "Matched"
                    ][
                        [
                            "Client_Internal_tracking",
                            "Goods_Description",
                            "Client_HS_code"
                        ]
                    ].copy()

                    st.dataframe(
                        client_preview,
                        use_container_width=True,
                        hide_index=True
                    )

                # Create the final download section.
                st.divider()

                st.subheader(
                    "📥 Final Output"
                )

                st.caption(
                    "Client_HS_code is collected from the third file. "
                    "The first file master determines the adjusted "
                    "HS_code. If Client_HS_code is not found in the "
                    "master, the original shipment HS_code remains unchanged."
                )

                st.download_button(
                    label="📥 Download Final Excel",
                    data=excel_file,
                    file_name=(
                        "APC_HS_Code_Validated_Output.xlsx"
                    ),
                    mime=(
                        "application/vnd.openxmlformats-"
                        "officedocument.spreadsheetml.sheet"
                    ),
                    use_container_width=True
                )

            except Exception as e:

                st.error(
                    "Failed to process the files."
                )

                st.exception(e)

    else:

        st.info(
            "Upload all three files to begin."
        )


if __name__ == "__main__":

    run()
